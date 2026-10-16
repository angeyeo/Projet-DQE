"""
Import d'un DQE Excel (.xlsx) dans la bibliothèque d'ouvrages du cabinet.

Deux temps, comme l'import de plan :
  1. analyser_dqe_excel() lit le fichier et renvoie un APERÇU : ouvrages
     candidats (désignation, unité, prix), lot reconnu, statut par rapport
     à la bibliothèque existante. Rien n'est écrit.
  2. L'utilisateur coche, corrige les lots, choisit un prix quand le fichier
     en contient plusieurs, puis confirme (vue OuvrageTypeViewSet.importer).

Ce que la lecture fait, et rien d'autre :
  - trouve la ligne d'en-tête (Désignation + Unité + Prix unitaire), dans
    n'importe quel ordre de colonnes, sur chaque feuille ;
  - lit les lignes qui portent une désignation et une unité ;
  - suit les titres de section pour reconnaître le lot par mots-clés ;
    un titre ambigu (« Gros œuvre » sans infra/super, « Menuiserie » seule)
    ne reçoit AUCUN lot : l'utilisateur le choisit ;
  - ignore sous-totaux, totaux, TVA et récapitulatifs ;
  - ne calcule rien : quantités et montants du fichier ne sont pas repris.
"""

import re
import unicodedata

TAILLE_MAX_OCTETS = 5 * 1024 * 1024
MAX_LIGNES = 5000

# Mots-clés -> lot. Ordre important : du plus précis au plus général.
MOTS_LOTS = (
    (("menuiserie bois", "menuiseries bois"), "lot_09_menuiserie_bois"),
    (("menuiserie metallique", "menuiseries metalliques", "menuiserie alu", "aluminium", "ferronnerie"),
     "lot_10_menuiserie_metallique"),
    (("faux plafond", "faux-plafond", "faux plafonds"), "lot_13_faux_plafond"),
    (("gros oeuvre - infrastructure", "infrastructure", "fondation", "fondations"),
     "lot_02_gros_oeuvre_infrastructure"),
    (("gros oeuvre - superstructure", "superstructure", "elevation", "elevations"),
     "lot_02_gros_oeuvre_superstructure"),
    (("terrassement", "terrassements"), "lot_01_terrassement"),
    (("etancheite",), "lot_03_etancheite"),
    (("plomberie", "sanitaire", "sanitaires"), "lot_04_plomberie"),
    (("assainissement",), "lot_05_assainissement"),
    (("electricite",), "lot_06_electricite"),
    (("charpente",), "lot_07_charpente"),
    (("couverture", "toiture"), "lot_08_couverture"),
    (("revetement", "revetements", "carrelage", "faience"), "lot_11_revetements"),
    (("peinture", "peintures"), "lot_12_peinture"),
    (("vrd", "voirie", "reseaux divers"), "lot_14_vrd"),
    (("generalites", "installation de chantier", "amenee et repli", "installation et repli"),
     "lot_00_generalites"),
    (("divers",), "lot_15_divers"),
)
# Titres reconnus comme un lot mais sans lot précis possible.
AMBIGUS = (("gros oeuvre", "maconnerie", "beton arme"), ("menuiserie", "menuiseries"))

MOTS_ENTETE = {
    "designation": ("designation", "designations", "libelle", "description", "designation des ouvrages"),
    "unite": ("unite", "unites", "u", "u.", "unit"),
    "prix": ("p.u", "pu", "p.u.", "prix unitaire", "prix unitaires", "prix u", "p.u ht", "pu ht",
             "prix unitaire (fcfa)", "p.u. (fcfa)", "prix unitaire ht"),
}
EXCLUS = re.compile(r"^(sous[- ]?total|total|tva|montant|recapitulatif|arrete|net a payer|report)", re.IGNORECASE)


class FichierDQEInvalide(Exception):
    pass


def _compter_formules(fichier, positions):
    """Parmi les prix vides, ceux qui sont des formules dont Excel n'a pas
    enregistré la valeur (fichier généré par un programme, jamais ouvert)."""
    if not positions:
        return 0
    import openpyxl

    if hasattr(fichier, "seek"):
        fichier.seek(0)
    classeur = openpyxl.load_workbook(fichier, data_only=False)
    n = sum(1 for f, r, c in positions
            if isinstance(classeur[f].cell(row=r, column=c).value, str)
            and classeur[f].cell(row=r, column=c).value.startswith("="))
    classeur.close()
    return n


def normaliser(texte):
    """Minuscules, sans accents (œ -> oe), espaces simples (comparaisons uniquement)."""
    t = str(texte or "").replace("œ", "oe").replace("Œ", "OE").replace("æ", "ae").replace("Æ", "AE")
    t = unicodedata.normalize("NFKD", t).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t).strip().lower()


_norm_oe = normaliser
# Un titre de section est court ; une longue phrase (hypothèse, note) n'en est pas un.
LONGUEUR_MAX_TITRE = 80


def reconnaitre_lot(texte):
    """(lot | None, ambigu: bool) d'après un titre de section ou de feuille."""
    t = " " + re.sub(r"[^a-z0-9 \-]", " ", _norm_oe(texte)) + " "
    t = re.sub(r"\s+", " ", t)
    for mots, lot in MOTS_LOTS:
        if any(f" {m} " in t or t.strip().startswith(m) for m in mots):
            return lot, False
    if any(any(f" {m} " in t for m in groupe) for groupe in AMBIGUS):
        return None, True
    return None, False


def lire_prix(valeur):
    """Nombre >= 0, ou None. Accepte « 85 000 », « 85 000 FCFA », « 1 250,5 »."""
    if valeur is None or isinstance(valeur, bool):
        return None
    if isinstance(valeur, (int, float)):
        return float(valeur) if valeur >= 0 else None
    t = re.sub(r"(?i)f\s?cfa|fcfa|f\b|cfa", "", str(valeur)).replace(" ", " ").replace("\xa0", " ").strip()
    t = t.replace(" ", "").replace(",", ".")
    try:
        v = float(t)
    except ValueError:
        return None
    return v if v >= 0 else None


def _colonnes(ligne):
    """Indices des colonnes Désignation / Unité / Prix si la ligne est un en-tête."""
    idx = {}
    for i, v in enumerate(ligne):
        t = _norm_oe(v).rstrip(" :")
        if not t:
            continue
        for cle, mots in MOTS_ENTETE.items():
            if cle in idx:
                continue
            if t in mots or (cle != "unite" and any(t.startswith(m) for m in mots if len(m) > 3)):
                idx[cle] = i
                break
    return idx if len(idx) == 3 else None


def analyser_dqe_excel(fichier, ouvrages_existants=()):
    """Aperçu de l'import. `fichier` : chemin ou objet fichier (.xlsx).
    `ouvrages_existants` : [(lot, designation, unite, prix_unitaire)] de la
    bibliothèque du cabinet, pour qualifier chaque ligne."""
    try:
        import openpyxl
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:  # pragma: no cover
        raise FichierDQEInvalide(f"Lecteur Excel indisponible : {exc}") from exc
    import zipfile

    try:
        classeur = openpyxl.load_workbook(fichier, data_only=True, read_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise FichierDQEInvalide(
            "Ce fichier n'est pas un classeur Excel .xlsx lisible. Un fichier .xls (ancien format) doit être "
            "enregistré en .xlsx depuis Excel.") from exc
    except Exception as exc:  # zip corrompu, etc.
        raise FichierDQEInvalide(f"Classeur Excel illisible : {exc}") from exc

    existants = {(lot, normaliser(d), normaliser(u)): p for lot, d, u, p in ouvrages_existants}
    candidats, ordre = {}, []
    feuilles, ignorees, avertissements = [], 0, []
    prix_vides = []

    for feuille in classeur.worksheets:
        lot_feuille, ambigu_feuille = reconnaitre_lot(feuille.title)
        contexte = {"lot": lot_feuille, "titre": feuille.title if (lot_feuille or ambigu_feuille) else None}
        colonnes, lues = None, 0
        for num, ligne in enumerate(feuille.iter_rows(values_only=True), start=1):
            if num > MAX_LIGNES:
                avertissements.append(f"Feuille « {feuille.title} » : lecture arrêtée à {MAX_LIGNES} lignes.")
                break
            if colonnes is None:
                colonnes = _colonnes(ligne)
                if colonnes is None:
                    titre = " ".join(str(v) for v in ligne if isinstance(v, str))
                    lot, ambigu = reconnaitre_lot(titre) if len(titre) <= LONGUEUR_MAX_TITRE else (None, False)
                    if lot or ambigu:
                        contexte = {"lot": lot, "titre": " ".join(str(v) for v in ligne if isinstance(v, str))}
                continue
            cell = lambda cle: ligne[colonnes[cle]] if colonnes[cle] < len(ligne) else None  # noqa: E731
            designation = str(cell("designation") or "").strip()
            unite = str(cell("unite") or "").strip()
            brut_prix = cell("prix")
            if not designation:
                continue
            if EXCLUS.match(_norm_oe(designation)):
                continue
            if not unite and lire_prix(brut_prix) is None:
                # Titre de section : change le lot seulement s'il en désigne un.
                if len(designation) > LONGUEUR_MAX_TITRE:
                    continue
                lot, ambigu = reconnaitre_lot(designation)
                if lot or ambigu or designation.upper().startswith("LOT"):
                    contexte = {"lot": lot, "titre": designation}
                continue
            if not unite:
                ignorees += 1
                continue
            prix = lire_prix(brut_prix)
            if prix is None:
                prix_vides.append((feuille.title, num, colonnes["prix"] + 1))
            lues += 1
            cle = (contexte["lot"], normaliser(designation), normaliser(unite))
            c = candidats.get(cle)
            if c is None:
                c = candidats[cle] = {
                    "lot": contexte["lot"], "section": contexte["titre"], "designation": designation[:200],
                    "unite": unite[:20], "prix_possibles": [], "occurrences": 0,
                    "origine": f"{feuille.title}, ligne {num}",
                }
                ordre.append(cle)
            c["occurrences"] += 1
            if prix is not None and prix not in c["prix_possibles"]:
                c["prix_possibles"].append(prix)
        if colonnes is not None:
            feuilles.append({"nom": feuille.title, "lignes": lues})
    classeur.close()
    formules_sans_valeur = _compter_formules(fichier, prix_vides)

    if not feuilles:
        raise FichierDQEInvalide(
            "Aucun tableau de DQE trouvé : le logiciel cherche une ligne d'en-tête contenant « Désignation », "
            "« Unité » et « Prix unitaire » (ou « P.U. »).")

    lignes = []
    for cle in ordre:
        c = candidats[cle]
        prix = c["prix_possibles"][0] if len(c["prix_possibles"]) == 1 else None
        if c["lot"] is None:
            statut = "lot_a_choisir"
        elif len(c["prix_possibles"]) > 1:
            statut = "prix_a_arbitrer"
        elif not c["prix_possibles"]:
            statut = "sans_prix"
        else:
            existant = existants.get((c["lot"], normaliser(c["designation"]), normaliser(c["unite"])), "absent")
            statut = ("nouveau" if existant == "absent" else "identique" if existant == prix else "prix_different")
            c["prix_bibliotheque"] = None if existant == "absent" else existant
        lignes.append({**c, "prix_unitaire": prix, "statut": statut,
                       "a_importer": statut in ("nouveau",)})

    if formules_sans_valeur:
        avertissements.append(
            f"{formules_sans_valeur} prix calculé(s) par une formule sans valeur enregistrée : ouvrez le fichier dans "
            "Excel, enregistrez-le, puis importez-le à nouveau.")
    if ignorees:
        avertissements.append(f"{ignorees} ligne(s) sans unité ignorée(s).")
    compte = {}
    for l in lignes:
        compte[l["statut"]] = compte.get(l["statut"], 0) + 1
    return {"feuilles": feuilles, "lignes": lignes, "compte": compte, "avertissements": avertissements}
