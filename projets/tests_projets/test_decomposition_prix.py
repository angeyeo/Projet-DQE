"""Prix matériaux / main-d'œuvre séparés (issue #92).

CAS DE RÉFÉRENCE recalculé à la main (prix DE TEST) :
  Porte isoplane : matériaux 60 000 + main-d'œuvre 25 000 = 85 000 FCFA/u
  poste de 3 u   : montant 3 × 85 000 = 255 000
                   matériaux 3 × 60 000 = 180 000 ; main-d'œuvre 255 000 − 180 000 = 75 000
  Peinture (non décomposée) : 2 500 FCFA/m² × 40 m² = 100 000 -> « non décomposé »
  Total 355 000 = 180 000 + 75 000 + 100 000 (non décomposé).

  Arrondi : 0,333 m³ × (matériaux 70 001 + MO 29 999) -> montant 33 300,
  matériaux round(0,333 × 70 001) = round(23 310,333) = 23 310, MO 33 300 − 23 310 = 9 990.
"""
import io

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from rest_framework.test import APITestCase

from projets.models import OuvrageType, Profil, Projet
from projets.services.decomposition_prix import decorer_ligne, totaux, verifier_decomposition
from projets.services.dqe_calculator import calculer_projet_dqe
from projets.services.dqe_exporters import exporter_dqe_excel, exporter_dqe_pdf
from projets.tests_projets.utils import BAREME_TEST, authentifier

URL = "/api/ouvrages/"


class Regles(SimpleTestCase):
    def test_verification(self):
        self.assertEqual(verifier_decomposition(85000, None, None), (85000, None, None))
        self.assertEqual(verifier_decomposition(None, 60000, 25000), (85000, 60000, 25000))
        self.assertEqual(verifier_decomposition(85000, 60000, 25000), (85000, 60000, 25000))
        self.assertEqual(verifier_decomposition(85000, 85000, 0), (85000, 85000, 0))
        for args in ((85000, 60000, None), (85000, None, 25000), (85000, 60000, 20000), (None, -1, 5),
                     (None, "abc", 5)):
            with self.assertRaises(ValueError):
                verifier_decomposition(*args)

    def test_ligne_et_arrondi(self):
        ligne = decorer_ligne({"quantite": 0.333, "prix_unitaire": 100000, "montant": 33300}, 70001, 29999)
        self.assertEqual((ligne["montant_materiaux"], ligne["montant_main_oeuvre"]), (23310, 9990))
        self.assertEqual(ligne["montant_materiaux"] + ligne["montant_main_oeuvre"], ligne["montant"])
        nd = decorer_ligne({"quantite": 1, "prix_unitaire": 10, "montant": 10}, None, None)
        self.assertEqual((nd["decomposition"], nd["montant_materiaux"]), ("non_decompose", None))
        faux = decorer_ligne({"quantite": 1, "prix_unitaire": 10, "montant": 10}, 3, 3)
        self.assertEqual(faux["decomposition"], "non_decompose")
        self.assertTrue(faux["decomposition_incoherente"])
        self.assertEqual(totaux([ligne, nd]), {"materiaux": 23310, "main_oeuvre": 9990, "non_decompose": 10})


class Bibliotheque(APITestCase):
    def setUp(self):
        self.user, self.cabinet = authentifier(self, role=Profil.Role.INGENIEUR, bareme=BAREME_TEST)
        self.projet = Projet.objects.create(nom="Villa", entreprise=self.cabinet)

    def creer(self, **kw):
        donnees = {"lot": "lot_09_menuiserie_bois", "designation": "Porte isoplane", "unite": "u",
                   "prix_materiaux": 60000, "prix_main_oeuvre": 25000} | kw
        return self.client.post(URL, donnees, format="json")

    def poste(self, ouvrage, quantite):
        r = self.client.post("/api/postes-complementaires/", {
            "projet": self.projet.id, "mode": "simple", "ouvrage": ouvrage, "quantite": quantite}, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        return r.data

    def test_total_egal_a_la_somme(self):
        r = self.creer()
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual((r.data["prix_unitaire"], r.data["prix_materiaux"], r.data["prix_main_oeuvre"]),
                         (85000, 60000, 25000))
        r = self.creer(designation="Autre", prix_unitaire=90000)  # total contradictoire
        self.assertEqual(r.status_code, 400)
        self.assertIn("n'est pas égal", str(r.data))
        r = self.creer(designation="Autre", prix_main_oeuvre=None, prix_unitaire=85000)  # une seule part
        self.assertEqual(r.status_code, 400)
        self.assertIn("DEUX parts", str(r.data))
        # Prix sans décomposition : toujours valable.
        r = self.creer(designation="Sans parts", prix_materiaux=None, prix_main_oeuvre=None, prix_unitaire=50000)
        self.assertEqual(r.status_code, 201, r.data)
        self.assertIsNone(r.data["prix_materiaux"])

    def test_dqe_cas_de_reference(self):
        porte = self.creer().data["id"]
        peinture = self.creer(lot="lot_12_peinture", designation="Peinture", unite="m²", prix_unitaire=2500,
                              prix_materiaux=None, prix_main_oeuvre=None).data["id"]
        p = self.poste(porte, 3)
        self.assertEqual((p["prix_materiaux"], p["prix_main_oeuvre"]), (60000, 25000))
        self.poste(peinture, 40)
        d = self.client.post(f"/api/projets/{self.projet.id}/generer-dqe/", format="json").json()
        l_porte = next(l for l in d["lignes"] if l["designation"] == "Porte isoplane")
        self.assertEqual((l_porte["montant"], l_porte["montant_materiaux"], l_porte["montant_main_oeuvre"]),
                         (255000, 180000, 75000))
        self.assertEqual(l_porte["decomposition"], "decompose")
        l_peint = next(l for l in d["lignes"] if l["designation"] == "Peinture")
        self.assertEqual((l_peint["decomposition"], l_peint["montant_materiaux"]), ("non_decompose", None))
        self.assertEqual(d["decomposition"], {"materiaux": 180000, "main_oeuvre": 75000, "non_decompose": 100000,
                                              "lignes_non_decomposees": 1})
        self.assertEqual(d["total_general"], 355000)
        lot9 = next(l for l in d["lots"] if l["lot"] == "lot_09_menuiserie_bois")
        self.assertEqual(lot9["decomposition"], {"materiaux": 180000, "main_oeuvre": 75000, "non_decompose": 0})

        # Exports : colonnes présentes, « non décomposé » écrit, aucune part inventée.
        wb = openpyxl.load_workbook(exporter_dqe_excel(d))
        valeurs = [c for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row]
        self.assertIn("dont matériaux (FCFA)", valeurs)
        self.assertIn("non décomposé", valeurs)
        self.assertIn(180000, valeurs)
        self.assertIn(75000, valeurs)
        self.assertTrue(exporter_dqe_pdf(d).getvalue().startswith(b"%PDF"))

    def test_poste_saisi_avec_parts(self):
        r = self.client.post("/api/postes-complementaires/", {
            "projet": self.projet.id, "mode": "simple", "lot": "lot_12_peinture", "designation": "Enduit",
            "unite": "m²", "quantite": 10, "prix_materiaux": 1000, "prix_main_oeuvre": 1500}, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(r.data["prix_unitaire"], 2500)
        r = self.client.post("/api/postes-complementaires/", {
            "projet": self.projet.id, "mode": "simple", "lot": "lot_12_peinture", "designation": "Enduit 2",
            "unite": "m²", "quantite": 10, "prix_unitaire": 3000, "prix_materiaux": 1000}, format="json")
        self.assertEqual(r.status_code, 400)
        self.assertIn("prix_materiaux", r.data)


class Bareme(APITestCase):
    def setUp(self):
        self.user, self.cabinet = authentifier(self, role=Profil.Role.ADMIN, bareme=BAREME_TEST)
        self.projet = Projet.objects.create(
            nom="Réf", entreprise=self.cabinet, usage_batiment="habitation", nb_niveaux=2, hauteur_etage=3.0,
            nb_travees_x=2, nb_travees_y=2, portee_x=4.0, portee_y=3.5, contrainte_sol_kn_m2=150)
        self.client.post(f"/api/projets/{self.projet.id}/generer_trame/")
        for e in self.projet.elements.all():
            self.client.post(f"/api/elements/{e.id}/valider/")

    def dqe(self):
        self.cabinet.refresh_from_db()
        return calculer_projet_dqe(self.projet, prix_unitaires=self.cabinet.get_prix_unitaires())

    def test_bareme_decompose(self):
        r = self.client.patch("/api/entreprise/", {
            "prix_decomposition": {"beton_m3": {"materiaux": 70000, "main_oeuvre": 30000}}}, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["prix_decomposition"], {"beton_m3": {"materiaux": 70000, "main_oeuvre": 30000}})
        d = self.dqe()
        beton = [l for l in d["lignes"] if l.get("cle_prix") == "beton_m3"]
        self.assertTrue(beton)
        for l in beton:
            self.assertEqual(l["decomposition"], "decompose")
            self.assertEqual(l["montant_materiaux"] + l["montant_main_oeuvre"], l["montant"])
        acier = next(l for l in d["lignes"] if l.get("cle_prix") == "acier_kg")
        self.assertEqual(acier["decomposition"], "non_decompose")
        dec = d["decomposition"]
        self.assertEqual(dec["materiaux"] + dec["main_oeuvre"] + dec["non_decompose"], d["total_general"])
        for lot in d["lots"]:
            t = lot["decomposition"]
            self.assertEqual(t["materiaux"] + t["main_oeuvre"] + t["non_decompose"], lot["sous_total"])

    def test_bareme_incoherent_refuse(self):
        r = self.client.patch("/api/entreprise/", {
            "prix_decomposition": {"beton_m3": {"materiaux": 70000, "main_oeuvre": 20000}}}, format="json")
        self.assertEqual(r.status_code, 400)
        self.assertIn("prix_decomposition", r.data)
        r = self.client.patch("/api/entreprise/", {
            "prix_decomposition": {"beton_m3": {"materiaux": 70000}}}, format="json")
        self.assertEqual(r.status_code, 400)

    def test_prix_change_sans_parts_refuse_puis_decomposition_retiree(self):
        self.client.patch("/api/entreprise/", {
            "prix_decomposition": {"beton_m3": {"materiaux": 70000, "main_oeuvre": 30000}}}, format="json")
        prix = dict(BAREME_TEST, beton_m3=110000)
        r = self.client.patch("/api/entreprise/", {"prix_unitaires": prix}, format="json")
        self.assertEqual(r.status_code, 400)  # jamais de parts devenues fausses en silence
        r = self.client.patch("/api/entreprise/", {"prix_unitaires": prix, "prix_decomposition": {}}, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["prix_decomposition"], {})


def classeur_parts():
    wb = openpyxl.Workbook()
    ws = wb.active
    for ligne in (
        ["N°", "Désignation", "U", "Qté", "P.U matériaux", "P.U main-d'œuvre", "P.U"],
        [None, "LOT 9 : MENUISERIE BOIS"],
        ["9.1", "Porte isoplane", "u", 3, 60000, 25000, 85000],
        ["9.2", "Fenêtre", "u", 2, 40000, 10000, None],          # total déduit = 50 000
        ["9.3", "Placard", "u", 1, 30000, 10000, 45000],         # incohérent -> non décomposé
        ["9.4", "Plinthe", "ml", 4, None, None, 1500],           # sans parts
    ):
        ws.append(ligne)
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


class ImportDeuxColonnes(APITestCase):
    def setUp(self):
        self.user, self.cabinet = authentifier(self, role=Profil.Role.INGENIEUR)

    def test_apercu_et_confirmation(self):
        f = SimpleUploadedFile("dqe.xlsx", classeur_parts())
        r = self.client.post(f"{URL}importer/", {"fichier": f}, format="multipart")
        self.assertEqual(r.status_code, 200, r.data)
        par = {l["designation"]: l for l in r.data["lignes"]}
        self.assertEqual((par["Porte isoplane"]["prix_unitaire"], par["Porte isoplane"]["prix_materiaux"],
                          par["Porte isoplane"]["prix_main_oeuvre"]), (85000, 60000, 25000))
        self.assertEqual((par["Fenêtre"]["prix_unitaire"], par["Fenêtre"]["prix_materiaux"]), (50000, 40000))
        self.assertEqual((par["Placard"]["prix_unitaire"], par["Placard"]["prix_materiaux"]), (45000, None))
        self.assertEqual((par["Plinthe"]["prix_unitaire"], par["Plinthe"]["prix_materiaux"]), (1500, None))
        self.assertTrue(any("ne fait pas le prix unitaire" in a for a in r.data["avertissements"]))

        lignes = [{k: par[n][k] for k in ("lot", "designation", "unite", "prix_unitaire", "prix_materiaux",
                                          "prix_main_oeuvre")} | {"action": "creer"} for n in par]
        r = self.client.post(f"{URL}importer/confirmer/", {"source_prix": "DQE test", "date_prix": "2026-10-23",
                                                  "lignes": lignes}, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        porte = OuvrageType.objects.get(designation="Porte isoplane")
        self.assertEqual((porte.prix_unitaire, porte.prix_materiaux, porte.prix_main_oeuvre), (85000, 60000, 25000))
        self.assertIsNone(OuvrageType.objects.get(designation="Placard").prix_materiaux)
