"""
Fiche d'approvisionnement : conversion des quantités du DQE en matériaux.

Cas de référence : volumes du DQE de l'ingénieur (R+1 NEW ERA), avec ses
corrections confirmées le 10/10/2026 :
  béton total = 31,474 + 29,643 + 6,923 = 68,04 m³
  acier = 6 213 kg ; agglos de 15 creux = 2 × 383,25 = 766,5 m²
  dosage 0,35 T / 0,45 m³ / 0,85 m³ par m³ ; pertes 3 % (agglos 4 %) ; 12 agglos/m².
"""

from io import BytesIO

from django.test import SimpleTestCase
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from projets.models import ElementStructurel, Profil, Projet
from projets.services.approvisionnement import calculer_fiche_approvisionnement, valider_parametres
from projets.tests_projets.utils import BAREME_TEST, authentifier


def dqe_reference():
    lignes = [{"repere": f"B{i}", "designation": d, "categorie": "BETON", "unite": "m³", "quantite": q}
              for i, (d, q) in enumerate([("Fondation", 31.474), ("RDC", 29.643), ("Étage", 6.923)])]
    lignes.append({"repere": "AC", "designation": "Armatures", "categorie": "ACIER", "unite": "kg", "quantite": 6213})
    lignes.append({"repere": "MA", "designation": "Agglos 15 creux (élévation)", "categorie": "MACONNERIE",
                   "unite": "m²", "quantite": 766.5, "cle_prix": "agglos_15_creux_m2"})
    lignes.append({"repere": "CO", "designation": "Coffrage", "categorie": "COFFRAGE", "unite": "m²", "quantite": 120})
    lignes.append({"repere": "EN", "designation": "Enduit", "categorie": "ENDUIT", "unite": "m²", "quantite": 500})
    return {"projet": {"nom": "R+1 NEW ERA"}, "lignes": lignes}


PRIX_INGENIEUR = {"ciment_t": 100000, "acier_t": 450000, "agglo_15_creux_u": 400}


class CalculFiche(SimpleTestCase):
    def fiche(self, params=None, prix=PRIX_INGENIEUR, tva=18):
        f = calculer_fiche_approvisionnement(dqe_reference(), params, prix, taux_tva_pct=tva)
        return f, {l["cle"]: l for l in f["lignes"]}

    def test_quantites_de_reference(self):
        f, l = self.fiche()
        self.assertEqual(f["volume_beton_m3"], 68.04)
        # 68,04 × 0,35 = 23,814 T ; × 1,03 = 24,52842 T -> 491 sacs de 50 kg
        self.assertEqual(l["ciment"]["quantite_nette"], 23.814)
        self.assertEqual(l["ciment"]["quantite_commande"], 24.528)
        self.assertEqual(l["ciment"]["conditionnement"], "491 sacs de 50 kg")
        # 68,04 × 0,45 × 1,03 = 31,53654 m³ ; 68,04 × 0,85 × 1,03 = 59,56902 m³
        self.assertEqual(l["sable"]["quantite_commande"], 31.537)
        self.assertEqual(l["gravier"]["quantite_commande"], 59.569)
        # 6 213 kg -> 6,213 T × 1,03 = 6,39939 T
        self.assertEqual(l["acier"]["quantite_commande"], 6.399)
        # 766,5 m² × 12 = 9 198 agglos ; × 1,04 = 9 565,92 -> 9 566
        self.assertEqual(l["agglo_15_creux_u"]["quantite_nette"], 9198)
        self.assertEqual(l["agglo_15_creux_u"]["quantite_commande"], 9566)
        self.assertEqual(l["agglo_15_creux_u"]["pertes_pct"], 4.0)

    def test_montants_et_total_partiel(self):
        f, l = self.fiche()
        # montant = quantité affichée (au millième) × prix, comme au DQE
        self.assertEqual(l["ciment"]["montant"], 2452800)       # 24,528 × 100 000
        self.assertEqual(l["acier"]["montant"], 2879550)        # 6,399 × 450 000
        self.assertEqual(l["agglo_15_creux_u"]["montant"], 3826400)  # 9 566 × 400
        self.assertIsNone(l["sable"]["montant"])
        self.assertIsNone(f["total_ht"])                        # sable et gravier sans prix
        self.assertEqual(f["total_chiffre"], 2452800 + 2879550 + 3826400)
        self.assertEqual({p["cle_prix"] for p in f["problemes"]}, {"sable_m3", "gravier_m3"})
        self.assertTrue(any(m.startswith("Total partiel") for m in f["messages"]))

    def test_total_complet_tva(self):
        prix = {**PRIX_INGENIEUR, "sable_m3": 10000, "gravier_m3": 15000}
        f, l = self.fiche(prix=prix)
        attendu = 2452800 + 315370 + 893535 + 2879550 + 3826400  # sable 31,537 × 10 000 ; gravier 59,569 × 15 000
        self.assertEqual(f["total_ht"], attendu)
        self.assertEqual(f["montant_tva"], round(attendu * 0.18))
        self.assertEqual(f["total_ttc"], attendu + f["montant_tva"])

    def test_voyages_seulement_si_capacite_saisie(self):
        f, l = self.fiche()
        self.assertIsNone(l["sable"]["voyages"]["nombre"])
        self.assertIn("non renseignée", l["sable"]["voyages"]["message"])
        f, l = self.fiche(params={"capacite_camion_sable_m3": 10})
        self.assertEqual(l["sable"]["voyages"]["nombre"], 4)   # 31,54 m³ / 10 m³
        self.assertEqual(f["parametres"]["capacite_camion_sable_m3"]["source"], "Paramètre du cabinet")

    def test_parametres_du_cabinet_et_sources(self):
        f, l = self.fiche(params={"pertes_pct": 5, "ciment_t_par_m3": 0.4})
        self.assertEqual(l["ciment"]["quantite_commande"], round(68.04 * 0.4 * 1.05, 3))
        self.assertEqual(f["parametres"]["pertes_agglos_pct"]["source"], "Ingénieur BTP, 10/10/2026")

    def test_lignes_non_converties_listees(self):
        f, _ = self.fiche()
        raisons = {n["designation"]: n["raison"] for n in f["non_convertis"]}
        self.assertIn("coffrage", raisons["Coffrage"])
        self.assertIn("dosage du mortier", raisons["Enduit"])

    def test_dqe_vide(self):
        f = calculer_fiche_approvisionnement({"lignes": []}, None, {}, taux_tva_pct=18)
        self.assertEqual(f["lignes"], [])
        self.assertIsNone(f["total_ht"])
        self.assertIn("Aucune quantité", f["messages"][0])

    def test_validation_des_parametres(self):
        self.assertEqual(valider_parametres({"pertes_pct": "3", "capacite_camion_sable_m3": ""}), {"pertes_pct": 3.0})
        for mauvais in ({"inconnu": 1}, {"pertes_pct": -1}, {"ciment_t_par_m3": 3.5}, {"agglos_par_m2": "abc"}):
            with self.subTest(mauvais=mauvais), self.assertRaises(ValueError):
                valider_parametres(mauvais)


class APIApprovisionnement(APITestCase):
    def setUp(self):
        self.user, self.cabinet = authentifier(self, role=Profil.Role.ADMIN,
                                               bareme={**BAREME_TEST, "ciment_t": 100000})
        self.projet = Projet.objects.create(nom="Appro", entreprise=self.cabinet, usage_batiment="habitation",
                                            nb_niveaux=1)
        ElementStructurel.objects.create(
            projet=self.projet, type_element=ElementStructurel.TypeElement.POTEAU, identifiant="P1",
            hauteur_poteau=3.0, resultat_calcul={"cote_cm": 20}, resultat_valide={"cote_cm": 20},
            statut=ElementStructurel.Statut.VALIDE)

    def test_fiche_json(self):
        r = self.client.get(f"/api/projets/{self.projet.id}/approvisionnement/")
        self.assertEqual(r.status_code, 200, r.data)
        ciment = next(l for l in r.data["lignes"] if l["cle"] == "ciment")
        # Poteau 0,2 × 0,2 × 3 = 0,12 m³ -> 0,12 × 0,35 × 1,03 = 0,04326 T
        self.assertEqual(ciment["quantite_commande"], 0.043)
        self.assertEqual(ciment["montant"], 4300)  # 0,043 × 100 000

    def test_export_excel(self):
        r = self.client.get(f"/api/projets/{self.projet.id}/approvisionnement/?export=excel")
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(BytesIO(r.content)).active
        self.assertEqual(ws["A2"].value, "FICHE D'APPROVISIONNEMENT EN MATÉRIAUX")
        self.assertEqual(self.client.get(f"/api/projets/{self.projet.id}/approvisionnement/?export=pdf").status_code, 400)

    def test_elements_non_valides_bloquent(self):
        self.projet.elements.update(statut=ElementStructurel.Statut.PROPOSE)
        r = self.client.get(f"/api/projets/{self.projet.id}/approvisionnement/")
        self.assertEqual(r.status_code, 400)

    def test_autre_cabinet_404(self):
        from projets.tests_projets.utils import creer_cabinet, creer_membre
        autre = creer_membre(creer_cabinet("Autre"), username="intrus")
        self.client.force_authenticate(user=autre)
        self.assertEqual(self.client.get(f"/api/projets/{self.projet.id}/approvisionnement/").status_code, 404)

    def test_parametres_cabinet_valides_par_l_api(self):
        r = self.client.patch("/api/entreprise/", {"parametres_approvisionnement": {"pertes_pct": 80}}, format="json")
        self.assertEqual(r.status_code, 400)
        r = self.client.patch("/api/entreprise/", {"parametres_approvisionnement": {"pertes_pct": 5}}, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        self.cabinet.refresh_from_db()
        self.assertEqual(self.cabinet.parametres_approvisionnement, {"pertes_pct": 5.0})
